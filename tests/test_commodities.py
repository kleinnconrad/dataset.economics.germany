"""Tests of the World Bank Pink Sheet collector with a generated workbook."""
import io

import openpyxl
import pytest

from collectors.commodities import PAGE_URL, CommoditiesCollector, WorkbookError, parse_sheet, parse_workbook
from conftest import CATALOG_FIXTURE, read_dataset
from state import CollectorEntry

WORKBOOK_URL = "https://thedocs.worldbank.org/en/doc/abc-0050012026/related/CMO-Historical-Data-Monthly.xlsx"


def workbook(brent_september: float = 116.8) -> bytes:
    """Builds a workbook with the layout of the Pink Sheet."""
    book = openpyxl.Workbook()
    prices = book.active
    prices.title = "Monthly Prices"
    prices.append(["World Bank Commodity Price Data (The Pink Sheet)"])
    prices.append(["Updated on October 02, 2026"])
    prices.append([None, "Crude oil, Brent", "Coal, South African **"])
    prices.append([None, "($/bbl)", "($/mt)"])
    prices.append(["2026M08", 110.5, "…"])
    prices.append(["2026M09", brent_september, 98.2])
    indices = book.create_sheet("Monthly Indices")
    indices.append(["monthly indices based on nominal US dollars, 2010=100"])
    indices.append([None, "Total Index", "Energy", "Non-energy **"])
    indices.append([None, None, None, " "])
    indices.append(["2026M09", 142.8, 148.9, 130.5])
    book.create_sheet("Description").append(["Series Description"])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def test_parse_sheet_joins_header_rows_and_reads_units() -> None:
    sheet = parse_sheet([
        ("title",), (None, "Crude oil, Brent", "Coal, South African **"), (None, "($/bbl)", "($/mt)"),
        ("1960M01", 1.6, "…"), ("1960M02", 1.7, 9.0),
    ])
    assert sheet.columns == {"Crude oil, Brent": {"1960-01": 1.6, "1960-02": 1.7},
                             "Coal, South African": {"1960-02": 9.0}}
    assert sheet.units["Crude oil, Brent"] == "$/bbl"
    with pytest.raises(WorkbookError):
        parse_sheet([("no periods",)])


def test_parse_workbook_reads_prices_indices_and_update_date() -> None:
    sheets, updated = parse_workbook(workbook())
    assert sheets["Monthly Prices"].columns["Crude oil, Brent"] == {"2026-08": 110.5, "2026-09": 116.8}
    assert sheets["Monthly Indices"].columns["Energy"] == {"2026-09": 148.9}
    assert "Non-energy" in sheets["Monthly Indices"].columns
    assert "Description" not in sheets
    assert updated is not None and updated.date().isoformat() == "2026-10-02"
    with pytest.raises(WorkbookError):
        parse_workbook(b"not a workbook")


def test_collector_finds_the_workbook_and_uses_conditional_requests(make_ctx, fake_web, tmp_path) -> None:
    fake_web.add(PAGE_URL, f'<a href="{WORKBOOK_URL}">Monthly prices</a>')
    xlsx = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    fake_web.add(WORKBOOK_URL, workbook(), content_type=xlsx, headers={"etag": '"v1"'})
    ctx = make_ctx()
    result = CommoditiesCollector(CATALOG_FIXTURE).collect(ctx, CollectorEntry(first_run=ctx.today))
    assert result.error is None and result.counters["new"] == 2
    assert result.cursor == {"url": WORKBOOK_URL, "etag": '"v1"', "last_modified": None}
    records = {record["series_id"]: record for record in read_dataset(tmp_path / "data", "economic_series")}
    assert records["brent_crude_oil_price.WLD"]["values"] == {"2026-08": 110.5, "2026-09": 116.8}
    assert records["commodity_price_index_energy.WLD"]["source_updated"] == "2026-10-02T00:00:00Z"

    # Unchanged workbook: the server answers 304 to the conditional request
    fake_web.add(WORKBOOK_URL, b"", status=304, content_type=xlsx)
    entry = CollectorEntry(first_run=ctx.today, cursor=result.cursor)
    unchanged = CommoditiesCollector(CATALOG_FIXTURE).collect(make_ctx(), entry)
    assert unchanged.error is None and unchanged.counters["workbook"] == "unchanged"


def test_collector_falls_back_to_the_stored_address(make_ctx, fake_web, tmp_path) -> None:
    fake_web.add(WORKBOOK_URL, workbook(brent_september=120.0),
                 content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    ctx = make_ctx()
    entry = CollectorEntry(first_run=ctx.today, cursor={"url": WORKBOOK_URL})
    result = CommoditiesCollector(CATALOG_FIXTURE).collect(ctx, entry)
    assert result.error is None and result.counters["new"] == 2


def test_collector_without_any_address_reports_an_error(make_ctx) -> None:
    ctx = make_ctx()
    result = CommoditiesCollector(CATALOG_FIXTURE).collect(ctx, CollectorEntry(first_run=ctx.today))
    assert result.error is not None and "no workbook address" in result.error
